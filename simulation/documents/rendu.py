"""Rendu des blocs de contenu en PDF (texte), DOCX, ou PDF scanné (image seule)."""

from __future__ import annotations

import io
import random
from datetime import date
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

import pypdfium2 as pdfium
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from simulation.documents.contenu import Bloc

MENTION = "Document fictif généré pour le projet JurisMind (données simulées)"

_styles = getSampleStyleSheet()
STYLES = {
    "entete": ParagraphStyle("entete", parent=_styles["Normal"], fontSize=9, leading=11),
    "titre": ParagraphStyle("titre", parent=_styles["Title"], fontSize=14, spaceBefore=12, spaceAfter=12),
    "section": ParagraphStyle("section", parent=_styles["Heading4"], spaceBefore=8, spaceAfter=4),
    "para": ParagraphStyle(
        "para", parent=_styles["Normal"], fontSize=10.5, leading=14, alignment=TA_JUSTIFY, spaceAfter=6
    ),
    "signature": ParagraphStyle("signature", parent=_styles["Normal"], fontSize=10.5, alignment=TA_RIGHT),
    "centre": ParagraphStyle("centre", parent=_styles["Normal"], fontSize=9, alignment=TA_CENTER),
}


def _pdf_flowables(blocs: list[Bloc]) -> list[Any]:
    elements: list[Any] = []
    for type_, valeur in blocs:
        if type_ == "entete":
            elements += [Paragraph(f"<b>{escape(valeur[0])}</b>", STYLES["entete"])]
            elements += [Paragraph(escape(ligne), STYLES["entete"]) for ligne in valeur[1:]]
            elements.append(Spacer(1, 0.5 * cm))
        elif type_ == "titre":
            elements.append(Paragraph(escape(valeur), STYLES["titre"]))
        elif type_ == "section":
            elements.append(Paragraph(escape(valeur), STYLES["section"]))
        elif type_ == "para":
            elements.append(Paragraph(escape(valeur), STYLES["para"]))
        elif type_ == "liste":
            elements.append(
                ListFlowable(
                    [ListItem(Paragraph(escape(x), STYLES["para"])) for x in valeur],
                    bulletType="bullet",
                )
            )
        elif type_ == "tableau":
            cellules = [[Paragraph(escape(str(x)), STYLES["para"]) for x in ligne] for ligne in valeur]
            table = Table(cellules, repeatRows=1, hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEEEEE")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            elements += [table, Spacer(1, 0.3 * cm)]
        elif type_ == "signature":
            elements.append(Spacer(1, 0.6 * cm))
            elements += [Paragraph(escape(ligne) or "&nbsp;", STYLES["signature"]) for ligne in valeur]
    return elements


def pdf(blocs: list[Bloc], titre: str) -> bytes:
    tampon = io.BytesIO()
    doc = SimpleDocTemplate(
        tampon,
        pagesize=A4,
        leftMargin=2.2 * cm,
        rightMargin=2.2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title=titre,
        author="Cabinet Téranga Avocats (fictif)",
        subject=MENTION,
        creator="JurisMind simulation",
    )

    def pied(canvas: Any, document: Any) -> None:
        canvas.setFont("Helvetica", 8)
        canvas.drawCentredString(A4[0] / 2, 1.2 * cm, f"Page {document.page}")

    doc.build(_pdf_flowables(blocs), onFirstPage=pied, onLaterPages=pied)
    return tampon.getvalue()


def docx(blocs: list[Bloc], titre: str) -> bytes:
    d = DocxDocument()
    d.core_properties.title = titre
    d.core_properties.comments = MENTION
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    for type_, valeur in blocs:
        if type_ == "entete":
            for i, ligne in enumerate(valeur):
                run = d.add_paragraph().add_run(ligne)
                run.bold = i == 0
                run.font.size = Pt(9)
        elif type_ == "titre":
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(valeur)
            run.bold = True
            run.font.size = Pt(14)
        elif type_ == "section":
            d.add_paragraph().add_run(valeur).bold = True
        elif type_ == "para":
            d.add_paragraph(valeur).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif type_ == "liste":
            for x in valeur:
                d.add_paragraph(x, style="List Bullet")
        elif type_ == "tableau":
            table = d.add_table(rows=len(valeur), cols=len(valeur[0]))
            table.style = "Table Grid"
            for i, ligne in enumerate(valeur):
                for j, x in enumerate(ligne):
                    table.cell(i, j).text = str(x)
        elif type_ == "signature":
            for ligne in valeur:
                d.add_paragraph(ligne).alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tampon = io.BytesIO()
    d.save(tampon)
    return tampon.getvalue()


def _tampon(image: Image.Image, rng: random.Random, jour: date) -> None:
    """Tampon « RECU LE » penché, comme sur un courrier arrivé au cabinet (police intégrée : ASCII)."""
    largeur = image.width * 2 // 5
    hauteur = largeur * 2 // 5
    calque = Image.new("L", (largeur, hauteur), 255)
    dessin = ImageDraw.Draw(calque)
    police = ImageFont.load_default(size=hauteur // 6)
    dessin.rounded_rectangle((4, 4, largeur - 4, hauteur - 4), radius=hauteur // 10, outline=80, width=5)
    dessin.text((hauteur // 5, hauteur // 6), "CABINET TERANGA", fill=80, font=police)
    dessin.text((hauteur // 5, hauteur // 2 + 5), f"RECU LE {jour:%d/%m/%Y}", fill=80, font=police)
    calque = calque.rotate(rng.uniform(-18, 18), expand=True, fillcolor=255)
    x = rng.randint(image.width // 3, image.width - calque.width - 20)
    y = rng.randint(60, image.height // 4)
    image.paste(
        Image.composite(
            calque,
            image.crop((x, y, x + calque.width, y + calque.height)),
            calque.point(lambda v: 255 if v < 200 else 0),
        ),
        (x, y),
    )


def scanner(pdf_texte: bytes, graine: int, jour: date) -> bytes:
    """Transforme un PDF texte en PDF « scanné » : images seules, penchées, bruitées."""
    rng = random.Random(graine)
    dpi = rng.choice([110, 130, 150])
    pages: list[Image.Image] = []
    source = pdfium.PdfDocument(pdf_texte)
    for i in range(len(source)):
        image = source[i].render(scale=dpi / 72, grayscale=True).to_pil().convert("L")
        if i == 0:
            _tampon(image, rng, jour)
        image = image.rotate(
            rng.uniform(-2.0, 2.0),
            expand=False,
            fillcolor=rng.randint(235, 250),
            resample=Image.Resampling.BICUBIC,
        )
        image = image.filter(ImageFilter.GaussianBlur(rng.uniform(0.3, 0.9)))
        bruit = Image.effect_noise(image.size, rng.uniform(20, 45))
        image = Image.blend(image, bruit, rng.uniform(0.05, 0.12))
        image = ImageEnhance.Contrast(image).enhance(rng.uniform(0.8, 1.2))
        image = ImageEnhance.Brightness(image).enhance(rng.uniform(0.9, 1.05))
        pages.append(image)
    source.close()
    sortie = io.BytesIO()
    pages[0].save(sortie, format="PDF", save_all=True, append_images=pages[1:], resolution=dpi, quality=70)
    return sortie.getvalue()


def ecrire(
    chemin: Path, blocs: list[Bloc], titre: str, format: str, scan: bool, graine: int, jour: date
) -> None:
    chemin.parent.mkdir(parents=True, exist_ok=True)
    if format == "docx":
        chemin.write_bytes(docx(blocs, titre))
        return
    contenu = pdf(blocs, titre)
    chemin.write_bytes(scanner(contenu, graine, jour) if scan else contenu)
