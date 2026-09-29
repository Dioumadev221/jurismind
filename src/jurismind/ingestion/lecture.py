"""Lecture du contenu des fichiers du cabinet : PDF, PDF scannés (OCR) et Word.

Un tiers des documents du cabinet sont des scans : des images, sans texte. On ne le sait
qu'après avoir essayé de lire le texte du PDF — s'il n'y en a presque pas, c'est un scan,
et il faut passer par la reconnaissance de caractères (OCR).
"""

from __future__ import annotations

import logging
import os
from dataclasses import dataclass
from pathlib import Path

import pypdfium2 as pdfium
import pytesseract
from docx import Document as DocumentWord
from PIL import Image, ImageOps

from jurismind.core.config import get_settings

logger = logging.getLogger(__name__)

# En dessous de ce nombre de caractères par page, on considère que le PDF est un scan.
SEUIL_PAGE_VIDE = 80
# L'OCR lit mieux une image nette et grande : 300 points par pouce est la valeur usuelle.
PPP_OCR = 300
LANGUE_OCR = "fra"


class OcrIndisponible(RuntimeError):
    """Tesseract n'est pas installé : les scans ne peuvent pas être lus."""


@dataclass
class Page:
    numero: int  # numérotation humaine : la première page est la page 1
    texte: str


@dataclass
class Lecture:
    """Le contenu d'un fichier, page par page, et la façon dont il a été obtenu."""

    pages: list[Page]
    ocr: bool = False

    @property
    def texte(self) -> str:
        return "\n\n".join(page.texte for page in self.pages if page.texte.strip())

    @property
    def vide(self) -> bool:
        return not self.texte.strip()


def _configurer_tesseract() -> None:
    """Indique à Tesseract où se trouvent son exécutable et ses langues."""
    reglages = get_settings()
    if reglages.tesseract_exe:
        pytesseract.pytesseract.tesseract_cmd = reglages.tesseract_exe
    dossier = Path(reglages.tessdata_dir) if reglages.tessdata_dir else None
    if dossier and (dossier / f"{LANGUE_OCR}.traineddata").exists():
        # Variable lue par Tesseract lui-même : elle supporte les chemins avec espaces,
        # contrairement aux options passées en ligne de commande par pytesseract.
        os.environ["TESSDATA_PREFIX"] = str(dossier.resolve())


def ocr_disponible() -> bool:
    _configurer_tesseract()
    try:
        pytesseract.get_tesseract_version()
    except (pytesseract.TesseractNotFoundError, OSError):
        return False
    return True


def lire_image(image: Image.Image) -> str:
    """Reconnaissance de caractères sur une page scannée."""
    _configurer_tesseract()
    # Niveaux de gris puis étirement du contraste : le tampon et le bruit gênent moins.
    preparee = ImageOps.autocontrast(image.convert("L"))
    try:
        return str(pytesseract.image_to_string(preparee, lang=LANGUE_OCR))
    except pytesseract.TesseractNotFoundError as erreur:
        raise OcrIndisponible(
            "Tesseract est introuvable : installez-le (pack français) ou renseignez TESSERACT_EXE"
        ) from erreur


def lire_pdf(chemin: Path, ocr_si_necessaire: bool = True) -> Lecture:
    pdf = pdfium.PdfDocument(chemin)
    try:
        pages = [
            Page(numero=index + 1, texte=pdf[index].get_textpage().get_text_range().strip())
            for index in range(len(pdf))
        ]
        caracteres = sum(len(page.texte) for page in pages)
        if not pages or caracteres >= SEUIL_PAGE_VIDE * len(pages) or not ocr_si_necessaire:
            return Lecture(pages=pages)

        logger.info(
            "%s : %s caractères pour %s pages, passage par l'OCR", chemin.name, caracteres, len(pages)
        )
        pages_ocr = [
            Page(
                numero=index + 1,
                texte=lire_image(pdf[index].render(scale=PPP_OCR / 72).to_pil()).strip(),
            )
            for index in range(len(pdf))
        ]
        return Lecture(pages=pages_ocr, ocr=True)
    finally:
        pdf.close()


def lire_docx(chemin: Path) -> Lecture:
    document = DocumentWord(str(chemin))
    morceaux = [paragraphe.text for paragraphe in document.paragraphs]
    # Les tableaux (décomptes, états de loyers, répartitions du capital) portent des
    # chiffres essentiels : sans eux, l'extraction passerait à côté des montants.
    for tableau in document.tables:
        for ligne in tableau.rows:
            cellules = [cellule.text.strip() for cellule in ligne.cells]
            if any(cellules):
                morceaux.append(" | ".join(cellules))
    texte = "\n".join(morceau for morceau in morceaux if morceau.strip())
    # Un fichier Word n'a pas de pages au sens du PDF : tout tient dans une « page 1 ».
    return Lecture(pages=[Page(numero=1, texte=texte)])


def lire(chemin: Path, ocr_si_necessaire: bool = True) -> Lecture:
    """Lit un fichier du serveur de documents du cabinet."""
    if not chemin.exists():
        raise FileNotFoundError(chemin)
    suffixe = chemin.suffix.lower()
    if suffixe == ".pdf":
        return lire_pdf(chemin, ocr_si_necessaire=ocr_si_necessaire)
    if suffixe == ".docx":
        return lire_docx(chemin)
    raise ValueError(f"Format non pris en charge : {chemin.name}")
